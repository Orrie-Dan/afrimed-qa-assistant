"""Convert report/report.md to report/report.docx (Word), keeping tables, figures and links.

Usage: python -P scripts/md_to_docx.py [input.md] [output.docx]

Handles the Markdown used in the report: #/##/### headings, paragraphs, nested - and 1. lists, pipe tables,
images, > quotes, --- rules, **bold**, *italic* / _italic_, `code`, [links](url), bare URLs and ^{superscript}.
Paragraphs under "References" get an APA hanging indent. Regenerate after editing the Markdown.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "report" / "report.md"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "report" / "report.docx"
LINK_COLOR = RGBColor(0x1F, 0x5F, 0xB4)

INLINE = re.compile(
    r"(?P<code>`[^`]+`)"
    r"|(?P<link>\[(?P<ltext>[^\]]+)\]\((?P<lurl>[^)]+)\))"
    r"|(?P<bold>\*\*(?P<btext>.+?)\*\*)"
    r"|(?P<ital>(?<![\w*])\*(?P<itext>[^*\s][^*]*?)\*(?![\w*])|(?<!\w)_(?P<utext>[^_\s][^_]*?)_(?!\w))"
    r"|(?P<sup>\^\{(?P<stext>[^}]+)\})"
    r"|(?P<url>https?://[^\s)|]+[^\s)|.,;:])"
)


def add_hyperlink(par, text, url, bold=False, italic=False, size=None):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                             is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FB4")
    under = OxmlElement("w:u")
    under.set(qn("w:val"), "single")
    rpr.extend([color, under])
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    par._p.append(link)


def add_inline(par, text, bold=False, italic=False, size=None):
    """Write Markdown inline text into a paragraph as formatted runs (recursive for nested emphasis)."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            r.bold, r.italic = bold or None, italic or None
            if size:
                r.font.size = Pt(size)
        if m.group("code"):
            r = par.add_run(m.group("code")[1:-1])
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt((size or 11) - 1)
            r.bold, r.italic = bold or None, italic or None
        elif m.group("link"):
            add_hyperlink(par, m.group("ltext"), m.group("lurl"), bold, italic, size)
        elif m.group("bold"):
            add_inline(par, m.group("btext"), True, italic, size)
        elif m.group("ital"):
            add_inline(par, m.group("itext") or m.group("utext"), bold, True, size)
        elif m.group("sup"):
            r = par.add_run(m.group("stext"))
            r.font.superscript = True
            r.bold, r.italic = bold or None, italic or None
        elif m.group("url"):
            add_hyperlink(par, m.group("url"), m.group("url"), bold, italic, size)
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        r.bold, r.italic = bold or None, italic or None
        if size:
            r.font.size = Pt(size)


def set_cell_shading(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def bottom_border(par, color="BFBFBF"):
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color)):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.append(bdr)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[2:]]
    ncol = len(header)
    has_header = any(header)               # "| | |" header = a plain two-column key/value table
    all_rows = ([header] if has_header else []) + body
    small = ncol >= 6
    table = doc.add_table(rows=len(all_rows), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # Column widths proportional to content length (markup stripped), with a floor so nothing gets crushed.
    plain = lambda t: re.sub(r"[*_`]|\[([^\]]*)\]\([^)]*\)", r"\1", t)
    need = [max(min(len(plain(r[j])) if j < len(r) else 0, 60) for r in all_rows) + 4 for j in range(ncol)]
    total_w = 6.27
    widths = [max(total_w * n / sum(need), 0.55) for n in need]
    scale = total_w / sum(widths)
    widths = [Inches(w * scale) for w in widths]
    for i, row in enumerate(all_rows):
        for j in range(ncol):
            cell = table.cell(i, j)
            cell.width = widths[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            is_head = has_header and i == 0
            add_inline(par, row[j] if j < len(row) else "", bold=is_head, size=8.5 if small else 9.5)
            if is_head:
                set_cell_shading(cell, "E8EEF7")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def restart_numbering(doc, par):
    """Give this numbered paragraph a fresh w:num (same list format, starting at 1) so lists don't continue."""
    style_numpr = doc.styles["List Number"].element.pPr.numPr
    numbering = doc.part.numbering_part.element
    old_num = style_numpr.numId.val
    abstract = next(n for n in numbering.findall(qn("w:num")) if n.numId == old_num).abstractNumId.val
    new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    absn = OxmlElement("w:abstractNumId")
    absn.set(qn("w:val"), str(abstract))
    num.append(absn)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_id


def set_num(par, num_id):
    ppr = par._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, nid])
    ppr.append(numpr)


def page_number_footer(section):
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tag)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._r.append(instr)


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)   # A4
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    page_number_footer(sec)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12.5)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    in_refs = False
    i = 0
    para_buf = []
    current_num = None   # numbering instance of the numbered list in progress (None = no list open)

    def flush():
        nonlocal para_buf
        if para_buf:
            p = doc.add_paragraph()
            add_inline(p, " ".join(para_buf))
            if in_refs:
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)
            para_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush(); i += 1; continue
        is_item = re.match(r"^\s*([-*]|\d+\.) ", line)
        if not is_item and line.startswith("  "):
            # Indented paragraph inside a list item: keep it under the item and keep the list's numbering.
            flush()
            text = []
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() \
                    and not re.match(r"^\s*([-*]|\d+\.) ", lines[i]):
                text.append(lines[i].strip()); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            add_inline(p, " ".join(text))
            continue
        if not is_item:
            current_num = None
        if stripped.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i]); i += 1
            add_table(doc, rows)
            continue
        m = re.match(r"^(#{1,3}) (.*)", line)
        if m:
            flush()
            level = len(m.group(1))
            text = m.group(2)
            if level == 1:
                p = doc.add_paragraph(style="Title")
                add_inline(p, text)
            else:
                p = doc.add_heading(level=level - 1)
                add_inline(p, text)
            in_refs = text.strip().endswith("References")
            i += 1
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if m:
            flush()
            img = (SRC.parent / m.group(2)).resolve()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img), width=Inches(6.2))
            i += 1
            continue
        if stripped == "---":
            flush()
            bottom_border(doc.add_paragraph())
            in_refs = False if not in_refs else in_refs
            i += 1
            continue
        if stripped.startswith(">"):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, stripped.lstrip("> ").strip(), italic=False)
            ppr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "8"), ("w:color", "9DB7DC")):
                left.set(qn(k), v)
            bdr.append(left)
            ppr.append(bdr)
            i += 1
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.) (.*)", line)
        if m:
            flush()
            depth = min(len(m.group(1)) // 2, 2)
            numbered = m.group(2)[0].isdigit()
            base = "List Number" if numbered else "List Bullet"
            style = base if depth == 0 else f"{base} {depth + 1}"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(3)
            if numbered and depth == 0:
                if current_num is None:          # first item of a new numbered list
                    current_num = restart_numbering(doc, p)
                set_num(p, current_num)
            text = m.group(3)
            i += 1
            # continuation lines indented under the item
            while i < len(lines) and lines[i].startswith(" " * (len(m.group(1)) + 2)) and lines[i].strip() \
                    and not re.match(r"^\s*([-*]|\d+\.) ", lines[i]):
                text += " " + lines[i].strip(); i += 1
            add_inline(p, text)
            continue
        para_buf.append(stripped)
        i += 1
    flush()
    doc.core_properties.title = "AfriMed Assistant"
    doc.core_properties.author = "Orrie-Dan"
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
